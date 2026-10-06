import os
import re
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional
from PIL import Image, ImageOps, ImageEnhance
try:
    import pymupdf as fitz  # PyMuPDF >= 1.24
except ImportError:  # pragma: no cover - legacy alias
    import fitz
import pytesseract
from app.core.config import settings

@dataclass
class ParsedDocument:
    text: str
    page_count: int
    metadata: Dict[str, Any] = field(default_factory=dict)
    tables: List[List[List[str]]] = field(default_factory=list)

def preprocess_image_for_ocr(image: Image.Image) -> Image.Image:
    """Grayscale and contrast enhancement for better OCR accuracy."""
    try:
        gray = image.convert("L")
        enhancer = ImageEnhance.Contrast(gray)
        enhanced = enhancer.enhance(1.8)
        return enhanced
    except Exception:
        return image

def run_tesseract_ocr(image: Image.Image) -> str:
    """Run Tesseract OCR on a PIL image with graceful fallback."""
    try:
        preprocessed = preprocess_image_for_ocr(image)
        text = pytesseract.image_to_string(preprocessed, lang=settings.OCR_LANG)
        return text.strip()
    except Exception:
        return ""

def _parse_pdf_text_fallback(file_path: str) -> ParsedDocument:
    """
    Last-resort fallback when the file cannot be opened as a valid PDF
    (e.g. corrupt or mislabeled upload). Attempts a plain-text decode so
    downstream regex extraction can still run instead of hard-failing.
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
        return ParsedDocument(text=text, page_count=1, metadata={"source": "pdf_text_fallback"})
    except Exception:
        return ParsedDocument(text="", page_count=0, metadata={"source": "pdf_unreadable"})

def parse_pdf(file_path: str) -> ParsedDocument:
    """
    Extracts text from PDF. If yield is < 50 characters per page,
    falls back to page-by-page OCR rendering. If the file is not a
    valid PDF at all, falls back to plain-text decoding.
    """
    try:
        doc = fitz.open(file_path)
    except Exception:
        return _parse_pdf_text_fallback(file_path)
    page_count = len(doc)
    extracted_text_pages = []
    ocr_needed = False

    for page_num in range(page_count):
        page = doc[page_num]
        raw_text = page.get_text()
        text = str(raw_text).strip() if isinstance(raw_text, str) else ""
        if len(text) < 50:
            ocr_needed = True
        extracted_text_pages.append(text)

    # If low yield detected across pages, attempt OCR fallback
    if ocr_needed or page_count == 0:
        ocr_text_pages = []
        for page_num in range(page_count):
            page = doc[page_num]
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            ocr_result = run_tesseract_ocr(img)
            # Use OCR if it yielded more text than native extraction
            page_text = ocr_result if len(ocr_result) > len(extracted_text_pages[page_num]) else extracted_text_pages[page_num]
            ocr_text_pages.append(page_text)
        full_text = "\n\n--- Page Break ---\n\n".join(ocr_text_pages)
    else:
        full_text = "\n\n--- Page Break ---\n\n".join(extracted_text_pages)

    doc.close()
    return ParsedDocument(text=full_text, page_count=page_count, metadata={"source": "pdf"})

def parse_image(file_path: str) -> ParsedDocument:
    """Extracts text from PNG, JPEG, or TIFF using Tesseract OCR."""
    with Image.open(file_path) as img:
        text = run_tesseract_ocr(img)
    return ParsedDocument(text=text, page_count=1, metadata={"source": "image"})

def parse_docx(file_path: str) -> ParsedDocument:
    """Extracts text from Word documents (.docx)."""
    try:
        import docx
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)
        return ParsedDocument(text="\n".join(paragraphs), page_count=1, metadata={"source": "docx"})
    except Exception:
        return ParsedDocument(text="", page_count=1, metadata={"source": "docx_error"})

def parse_xlsx(file_path: str) -> ParsedDocument:
    """Extracts text and table rows from Excel spreadsheets (.xlsx)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            lines.append(f"Sheet: {sheet_name}")
            if hasattr(sheet, "iter_rows"):
                for row in sheet.iter_rows(values_only=True):  # type: ignore
                    non_empty = [str(val).strip() for val in row if val is not None and str(val).strip()]
                    if non_empty:
                        lines.append(" | ".join(non_empty))
        return ParsedDocument(text="\n".join(lines), page_count=len(wb.sheetnames), metadata={"source": "xlsx"})
    except Exception:
        return ParsedDocument(text="", page_count=1, metadata={"source": "xlsx_error"})

def parse_document(file_path: str, mime_type: Optional[str] = None) -> ParsedDocument:
    """Dispatches document to the appropriate parser based on extension or MIME type."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf" or mime_type == "application/pdf":
        return parse_pdf(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".tif"] or (mime_type and mime_type.startswith("image/")):
        return parse_image(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    elif ext in [".xlsx", ".xls"]:
        return parse_xlsx(file_path)
    else:
        # Fallback to plain text read if possible
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return ParsedDocument(text=f.read(), page_count=1, metadata={"source": "raw_text"})
        except Exception:
            return ParsedDocument(text="", page_count=0, metadata={"source": "unsupported"})
