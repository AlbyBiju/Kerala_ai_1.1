"""Sub-Task 5 tests: parser dispatcher with real PDF fixtures."""
from app.utils.parsers import parse_document, parse_pdf


def test_parse_pdf_fixture_extracts_text(invoice_pdf_path):
    doc = parse_pdf(invoice_pdf_path)
    assert doc.page_count >= 1
    assert "COMMERCIAL INVOICE" in doc.text
    assert "1200" in doc.text
    assert "Apex Global Agro Exports Ltd." in doc.text


def test_parse_dispatcher_routes_pdf(invoice_pdf_path):
    doc = parse_document(invoice_pdf_path, "application/pdf")
    assert doc.metadata.get("source") == "pdf"
    assert "HS Code" in doc.text


def test_parse_corrupt_pdf_falls_back_to_text(tmp_path):
    fake = tmp_path / "broken.pdf"
    fake.write_bytes(b"this is not really a pdf, just plain text")
    doc = parse_document(str(fake), "application/pdf")
    # Must not raise; falls back to plain-text decode
    assert "not really a pdf" in doc.text


def test_parse_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Field", "Value"])
    ws.append(["Total Quantity", 500])
    path = tmp_path / "sheet.xlsx"
    wb.save(path)

    doc = parse_document(str(path), None)
    assert doc.metadata.get("source") == "xlsx"
    assert "Total Quantity" in doc.text
    assert "500" in doc.text


def test_parse_docx(tmp_path):
    import docx as docx_module

    d = docx_module.Document()
    d.add_paragraph("COMMERCIAL INVOICE")
    d.add_paragraph("Total Quantity: 900 Cartons")
    path = tmp_path / "doc.docx"
    d.save(str(path))

    doc = parse_document(str(path), None)
    assert doc.metadata.get("source") == "docx"
    assert "COMMERCIAL INVOICE" in doc.text
